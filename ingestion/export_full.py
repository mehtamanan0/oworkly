#!/usr/bin/env python3
"""Dump the FULL parsed workbook (all worker rows, not just the 10-row sample
`parse_wtg_skill_matrix.py`'s CLI prints) to JSON for the Node seeder."""
import json
import sys
from dataclasses import asdict

import openpyxl
from parse_wtg_skill_matrix import (
    unmerge_forward_fill,
    parse_header_block,
    parse_process_definitions,
    parse_worker_rows,
)

xlsx_path = sys.argv[1]
out_path = sys.argv[2]

wb = openpyxl.load_workbook(xlsx_path, data_only=True)
result = {"sheets": []}

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    grid = unmerge_forward_fill(ws)
    header = parse_header_block(grid, ws.max_column)
    processes = parse_process_definitions(grid, ws.max_column)
    workers = parse_worker_rows(grid, ws.max_row, ws.max_column, processes)
    result["sheets"].append({
        "sheet_name": sheet_name,
        "header": header,
        "processes": [asdict(p) for p in processes],
        "workers": [asdict(w) for w in workers],
    })
    print(f"{sheet_name}: {len(processes)} processes, {len(workers)} worker rows", file=sys.stderr)

with open(out_path, "w") as f:
    json.dump(result, f, indent=2, default=str)
print(f"Wrote {out_path}", file=sys.stderr)
